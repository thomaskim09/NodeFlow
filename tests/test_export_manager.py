import logging
import json
from pathlib import Path

import openpyxl

import database
from managers import export_manager
from repositories.base import initialize_database
from services.export_service import export_service


def _build_export_fixture():
    initialize_database()
    database.add_project("Exports")
    project_id = database.get_all_projects()[0]["id"]

    r1_id = database.add_participant(project_id, "R1")
    r2_id = database.add_participant(project_id, "R2")

    doc1_id = database.add_document(project_id, "Doc 1", "alpha beta", r1_id)
    doc2_id = database.add_document(project_id, "Doc 2", "gamma delta", r2_id)
    doc3_id = database.add_document(project_id, "Doc 3", "epsilon", r1_id)

    findings_a_id = database.add_node(project_id, "Findings A", None, "#111111")
    node_a1_id = database.add_node(project_id, "Node A1", findings_a_id, "#222222")
    sub_a1a_id = database.add_node(project_id, "Sub A1a", node_a1_id, "#333333")
    leaf_a1a1_id = database.add_node(project_id, "Leaf A1a1", sub_a1a_id, "#444444")
    findings_b_id = database.add_node(project_id, "Findings B", None, "#555555")

    database.add_coded_segment(doc1_id, findings_a_id, r1_id, 0, 5, "alpha")
    beta_id = database.add_coded_segment(doc1_id, node_a1_id, r1_id, 6, 10, "beta")
    gamma_id = database.add_coded_segment(doc2_id, sub_a1a_id, r2_id, 0, 5, "gamma")
    database.add_coded_segment(doc2_id, leaf_a1a1_id, r2_id, 6, 11, "delta")
    database.add_coded_segment(doc3_id, findings_b_id, r1_id, 0, 7, "epsilon")
    database.update_coded_segment_remark(beta_id, "Beta remark")
    database.update_coded_segment_remark(gamma_id, "Gamma remark")

    return {
        "project_id": project_id,
        "r1_id": r1_id,
        "r2_id": r2_id,
        "findings_a_id": findings_a_id,
        "findings_b_id": findings_b_id,
    }


def test_classification_workbook_exports_quotes_and_matrix(tmp_path, monkeypatch):
    fixture = _build_export_fixture()
    export_path = tmp_path / "classification.xlsx"

    monkeypatch.setattr(
        export_manager,
        "_get_save_file_path",
        lambda *args, **kwargs: str(export_path),
    )
    monkeypatch.setattr(export_manager, "_show_export_saved", lambda *args: None)

    export_manager.export_classification_workbook(
        fixture["project_id"], language="English"
    )

    workbook = openpyxl.load_workbook(export_path)
    quotes_ws = workbook["Classification Quotes"]
    matrix_ws = workbook["Respondent Matrix"]

    assert [cell.value for cell in quotes_ws[1]] == [
        "Findings",
        "Node",
        "Sub-node",
        "Respondent",
        "Original quote",
        "Document",
        "Remark",
    ]
    assert list(quotes_ws.iter_rows(min_row=2, values_only=True)) == [
        ("Findings A", None, None, "R1", "alpha", "Doc 1", None),
        ("Findings A", "Node A1", None, "R1", "beta", "Doc 1", "Beta remark"),
        ("Findings A", "Node A1", "Sub A1a", "R2", "gamma", "Doc 2", "Gamma remark"),
        ("Findings A", "Node A1", "Sub A1a > Leaf A1a1", "R2", "delta", "Doc 2", None),
        ("Findings B", None, None, "R1", "epsilon", "Doc 3", None),
    ]

    assert [cell.value for cell in matrix_ws[1]] == [
        "Respondent",
        "Findings A",
        "Findings A > Node A1",
        "Findings A > Node A1 > Sub A1a",
        "Findings A > Node A1 > Sub A1a > Leaf A1a1",
        "Findings B",
    ]
    assert list(matrix_ws.iter_rows(min_row=2, values_only=True)) == [
        ("R1", 1, 1, 0, 0, 1),
        ("R2", 0, 0, 1, 1, 0),
    ]


def test_classification_workbook_respects_subtree_and_participant_filter(
    tmp_path, monkeypatch
):
    fixture = _build_export_fixture()
    export_path = tmp_path / "classification_filtered.xlsx"

    monkeypatch.setattr(
        export_manager,
        "_get_save_file_path",
        lambda *args, **kwargs: str(export_path),
    )
    monkeypatch.setattr(export_manager, "_show_export_saved", lambda *args: None)

    export_manager.export_classification_workbook(
        fixture["project_id"],
        participant_id=fixture["r2_id"],
        start_node_id=fixture["findings_a_id"],
        language="English",
    )

    workbook = openpyxl.load_workbook(export_path)
    quotes_ws = workbook["Classification Quotes"]
    matrix_ws = workbook["Respondent Matrix"]

    assert list(quotes_ws.iter_rows(min_row=2, values_only=True)) == [
        ("Findings A", "Node A1", "Sub A1a", "R2", "gamma", "Doc 2", "Gamma remark"),
        ("Findings A", "Node A1", "Sub A1a > Leaf A1a1", "R2", "delta", "Doc 2", None),
    ]
    assert [cell.value for cell in matrix_ws[1]] == [
        "Respondent",
        "Findings A",
        "Findings A > Node A1",
        "Findings A > Node A1 > Sub A1a",
        "Findings A > Node A1 > Sub A1a > Leaf A1a1",
    ]
    assert list(matrix_ws.iter_rows(min_row=2, values_only=True)) == [
        ("R2", 0, 0, 1, 1),
    ]
    assert "Findings B" not in [cell.value for cell in matrix_ws[1]]


def test_json_and_single_sheet_exports_include_remark(tmp_path, monkeypatch):
    fixture = _build_export_fixture()
    json_path = tmp_path / "segments.json"
    xlsx_path = tmp_path / "segments.xlsx"
    save_paths = iter([str(xlsx_path)])

    monkeypatch.setattr(
        export_manager,
        "_get_save_file_path",
        lambda *args, **kwargs: next(save_paths),
    )
    monkeypatch.setattr(
        export_manager.QFileDialog,
        "getSaveFileName",
        lambda *args, **kwargs: (str(json_path), "JSON Files (*.json)"),
    )
    monkeypatch.setattr(export_manager, "_show_export_saved", lambda *args: None)

    export_manager.export_to_json(fixture["project_id"])
    export_manager.export_project_to_excel_single_sheet(fixture["project_id"])

    exported = json.loads(json_path.read_text(encoding="utf-8"))
    beta_segment = exported[0]["children"][0]["segments"][0]
    assert beta_segment["remark"] == "Beta remark"

    workbook = openpyxl.load_workbook(xlsx_path)
    ws = workbook["Report"]
    assert [cell.value for cell in ws[1]] == [
        "Node",
        "Participant",
        "Coded Segment",
        "Document",
        "Remark",
    ]
    assert any(row[-1] == "Beta remark" for row in ws.iter_rows(min_row=2, values_only=True))


def test_get_save_file_path_logs_selected_path(caplog, monkeypatch):
    class FakeDialog:
        AcceptMode = type("AcceptMode", (), {"AcceptSave": object()})
        FileMode = type("FileMode", (), {"AnyFile": object()})
        Option = type(
            "Option",
            (),
            {"DontConfirmOverwrite": object(), "DontUseNativeDialog": object()},
        )
        DialogCode = type("DialogCode", (), {"Accepted": 1})

        def __init__(self, parent, title):
            self.title = title
            self.selected = ["/tmp/export.xlsx"]

        def setAcceptMode(self, *_args):
            pass

        def setFileMode(self, *_args):
            pass

        def setNameFilter(self, *_args):
            pass

        def setOption(self, *_args):
            pass

        def setDefaultSuffix(self, *_args):
            pass

        def setDirectory(self, *_args):
            pass

        def selectFile(self, *_args):
            pass

        def exec(self):
            return self.DialogCode.Accepted

        def selectedFiles(self):
            return self.selected

    monkeypatch.setattr(export_manager, "QFileDialog", FakeDialog)

    with caplog.at_level(logging.INFO):
        selected = export_manager._get_save_file_path(
            None, "Save Excel Report", "demo.xlsx", "Excel Files (*.xlsx)"
        )

    assert selected == "/tmp/export.xlsx"
    assert "Save dialog selected path: /tmp/export.xlsx" in caplog.text


def test_get_save_file_path_logs_cancel(caplog, monkeypatch):
    class FakeDialog:
        AcceptMode = type("AcceptMode", (), {"AcceptSave": object()})
        FileMode = type("FileMode", (), {"AnyFile": object()})
        Option = type(
            "Option",
            (),
            {"DontConfirmOverwrite": object(), "DontUseNativeDialog": object()},
        )
        DialogCode = type("DialogCode", (), {"Accepted": 1})

        def __init__(self, parent, title):
            self.title = title

        def setAcceptMode(self, *_args):
            pass

        def setFileMode(self, *_args):
            pass

        def setNameFilter(self, *_args):
            pass

        def setOption(self, *_args):
            pass

        def setDefaultSuffix(self, *_args):
            pass

        def setDirectory(self, *_args):
            pass

        def selectFile(self, *_args):
            pass

        def exec(self):
            return 0

        def selectedFiles(self):
            return []

    monkeypatch.setattr(export_manager, "QFileDialog", FakeDialog)

    with caplog.at_level(logging.INFO):
        selected = export_manager._get_save_file_path(
            None, "Save Excel Report", "demo.xlsx", "Excel Files (*.xlsx)"
        )

    assert selected == ""
    assert "Save dialog canceled" in caplog.text


def test_save_workbook_with_feedback_logs_and_surfaces_failure(caplog, monkeypatch):
    calls = []

    class BrokenWorkbook:
        def save(self, _path):
            raise RuntimeError("save failed")

    monkeypatch.setattr(
        export_manager.export_service,
        "show_unexpected_error",
        lambda parent, title, error: calls.append((parent, title, str(error))),
    )

    with caplog.at_level(logging.ERROR):
        export_manager._save_workbook_with_feedback(
            BrokenWorkbook(),
            "/tmp/export.xlsx",
            None,
            "Excel report successfully saved to:",
        )

    assert calls == [(None, "Export Error", "save failed")]
    assert "Workbook save failed for /tmp/export.xlsx" in caplog.text


def test_export_service_error_dialog_includes_log_path_and_details(monkeypatch):
    created = []

    class FakeMessageBox:
        Icon = type("Icon", (), {"Critical": object()})

        def __init__(self, parent):
            self.parent = parent
            self.icon = None
            self.window_title = ""
            self.text = ""
            self.informative_text = ""
            self.detailed_text = ""
            created.append(self)

        def setIcon(self, icon):
            self.icon = icon

        def setWindowTitle(self, title):
            self.window_title = title

        def setText(self, text):
            self.text = text

        def setInformativeText(self, text):
            self.informative_text = text

        def setDetailedText(self, text):
            self.detailed_text = text

        def exec(self):
            return 0

    monkeypatch.setattr("services.export_service.QMessageBox", FakeMessageBox)
    monkeypatch.setattr("services.export_service.get_log_path", lambda: Path("/tmp/nodeflow.log"))

    export_service.show_unexpected_error(
        None,
        "Export Error",
        RuntimeError("boom"),
        details="traceback details",
    )

    assert created
    dialog = created[0]
    assert dialog.window_title == "Export Error"
    assert "/tmp/nodeflow.log" in dialog.informative_text
    assert dialog.detailed_text == "traceback details"
