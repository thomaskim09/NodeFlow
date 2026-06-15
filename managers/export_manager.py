# managers/export_manager.py

import json
import logging
import re
from collections import Counter, defaultdict
from pathlib import Path

import database
import networkx as nx
import openpyxl
from PySide6.QtWidgets import QFileDialog, QMessageBox
from docx import Document
from docx.shared import RGBColor
from openpyxl.styles import Alignment, Font

from services.export_service import export_service
from services.platform_service import platform_service
from utils.common import get_translation

LOGGER = logging.getLogger(__name__)


def _get_project_name(project_id: int) -> str:
    for project in database.get_all_projects():
        if project["id"] == project_id:
            return project["name"]
    return f"project_{project_id}"


def _sanitize_file_stem(name: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]+', "_", name).strip()
    return cleaned or "NodeFlow_Export"


def _default_export_filename(stem: str, extension: str) -> str:
    return f"{_sanitize_file_stem(stem)}.{extension}"


def _get_participants_map(project_id: int) -> dict[int, str]:
    return {
        participant["id"]: participant["name"]
        for participant in database.get_participants_for_project(project_id)
    }


def _filter_segments_for_participant(coded_segments, participant_id):
    if participant_id is None:
        return coded_segments
    return [
        segment
        for segment in coded_segments
        if segment.get("participant_id") == participant_id
    ]


def _participant_file_suffix(project_id: int, participant_id) -> str:
    if participant_id is None:
        return "all_participants"
    participant_name = _get_participants_map(project_id).get(participant_id)
    return _sanitize_file_stem(participant_name or f"participant_{participant_id}")


def _show_export_saved(parent_widget, message, file_path):
    msg_box = QMessageBox(parent_widget)
    msg_box.setWindowTitle("Export Successful")
    msg_box.setText(f"{message}\n{file_path}")
    open_button = msg_box.addButton("Open File", QMessageBox.ActionRole)
    msg_box.addButton(QMessageBox.Ok)
    msg_box.exec()
    if msg_box.clickedButton() == open_button:
        platform_service.open_path(file_path)


def _sanitize_sheet_name(name: str) -> str:
    return re.sub(r"[\\/*?:\[\]]", "", name)[:31]


def _set_header_row(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)


def _set_column_widths(ws, widths: dict[str, int]):
    for column_letter, width in widths.items():
        ws.column_dimensions[column_letter].width = width


def _build_scoped_node_paths(nodes, start_node_id=None):
    nodes_by_id = {node["id"]: node for node in nodes}
    if start_node_id is not None and start_node_id not in nodes_by_id:
        return [], {}

    nodes_by_parent = {node_id: [] for node_id in nodes_by_id}
    nodes_by_parent[None] = []
    for node in nodes:
        nodes_by_parent.setdefault(node["parent_id"], []).append(node)
    for children in nodes_by_parent.values():
        children.sort(key=lambda item: item["position"])

    ordered_nodes = []
    path_by_id = {}

    def visit(node, current_path):
        ordered_nodes.append(node)
        path_by_id[node["id"]] = current_path
        for child in nodes_by_parent.get(node["id"], []):
            visit(child, current_path + [child["name"]])

    if start_node_id is None:
        for root in nodes_by_parent.get(None, []):
            visit(root, [root["name"]])
    else:
        start_node = nodes_by_id[start_node_id]
        visit(start_node, [start_node["name"]])

    return ordered_nodes, path_by_id


def _build_classification_columns(path_parts):
    findings = path_parts[0] if len(path_parts) >= 1 else ""
    node = path_parts[1] if len(path_parts) >= 2 else ""
    sub_node = " > ".join(path_parts[2:]) if len(path_parts) >= 3 else ""
    return findings, node, sub_node


def _participant_display_name(segment):
    return segment.get("participant_name") or "N/A"


def _log_export_start(export_kind: str, project_id, participant_id=None, node_id=None):
    LOGGER.info(
        "Starting %s export (project_id=%s, participant_id=%s, node_id=%s)",
        export_kind,
        project_id,
        participant_id,
        node_id,
    )


def _save_workbook_with_feedback(
    wb, file_path, parent_widget, success_message, error_title="Export Error"
):
    LOGGER.debug("Saving workbook to %s", file_path)
    try:
        wb.save(file_path)
        LOGGER.info("Workbook saved successfully: %s", file_path)
        _show_export_saved(parent_widget, success_message, file_path)
    except PermissionError:
        LOGGER.warning("Workbook save permission denied: %s", file_path)
        export_service.show_permission_error(parent_widget, file_path)
    except Exception as e:
        LOGGER.exception("Workbook save failed for %s", file_path)
        export_service.show_unexpected_error(parent_widget, error_title, e)


def _get_save_file_path(
    parent_widget,
    title: str,
    default_filename: str,
    file_filter: str,
) -> str:
    LOGGER.debug(
        "Opening save dialog title=%r default_filename=%r filter=%r",
        title,
        default_filename,
        file_filter,
    )
    dialog = QFileDialog(parent_widget, title)
    dialog.setAcceptMode(QFileDialog.AcceptMode.AcceptSave)
    dialog.setFileMode(QFileDialog.FileMode.AnyFile)
    dialog.setNameFilter(file_filter)
    dialog.setOption(QFileDialog.Option.DontConfirmOverwrite, False)
    dialog.setOption(QFileDialog.Option.DontUseNativeDialog, False)
    dialog.setDefaultSuffix(Path(default_filename).suffix.lstrip("."))

    default_path = Path.home() / default_filename if default_filename else Path.home()
    dialog.setDirectory(str(default_path.parent))
    dialog.selectFile(default_path.name)

    if dialog.exec() != QFileDialog.DialogCode.Accepted:
        LOGGER.info("Save dialog canceled for title=%r", title)
        return ""

    selected_files = dialog.selectedFiles()
    selected_path = selected_files[0] if selected_files else ""
    if selected_path:
        LOGGER.info("Save dialog selected path: %s", selected_path)
    else:
        LOGGER.warning("Save dialog accepted without a selected path for title=%r", title)
    return selected_path


def export_to_word(project_id, parent_widget=None):
    """Exports the coded segments of a project to a .docx file."""
    file_path, _ = QFileDialog.getSaveFileName(
        parent_widget, "Save Word Report", "", "Word Documents (*.docx)"
    )
    if not file_path:
        return

    # --- Fetch and structure the data ---
    nodes = database.get_nodes_for_project(project_id)
    coded_segments = database.get_coded_segments_for_project(project_id)

    # Group segments by their node ID
    segments_by_node = {}
    for seg in coded_segments:
        node_id = seg["node_id"]
        if node_id not in segments_by_node:
            segments_by_node[node_id] = []
        segments_by_node[node_id].append(seg)

    # --- Build the Word Document ---
    doc = Document()
    doc.add_heading("Qualitative Analysis Report", 0)

    def write_nodes_recursively(parent_id=None, level=0, prefix=""):
        children = sorted(
            [n for n in nodes if n["parent_id"] == parent_id],
            key=lambda x: x["position"],
        )

        for i, node in enumerate(children):
            current_prefix = f"{prefix}{i + 1}."
            doc.add_heading(f"{current_prefix} {node['name']}", level=level + 1)

            if node["id"] in segments_by_node:
                for segment_data in segments_by_node[node["id"]]:
                    participant = segment_data["participant_name"] or "N/A"
                    text = segment_data["content_preview"]
                    p = doc.add_paragraph(style="List Bullet")
                    p.add_run(f"{participant}: ").bold = True
                    p.add_run(text)

            write_nodes_recursively(
                parent_id=node["id"], level=level + 1, prefix=current_prefix
            )

    write_nodes_recursively()

    # --- Save the document with error handling ---
    try:
        doc.save(file_path)
        _show_export_saved(parent_widget, "Report successfully saved to:", file_path)
    except PermissionError:
        export_service.show_permission_error(parent_widget, file_path)
    except Exception as e:
        export_service.show_unexpected_error(parent_widget, "Export Error", e)


def export_to_json(project_id, parent_widget=None):
    """Exports the coded segments of a project to a .json file."""
    file_path, _ = QFileDialog.getSaveFileName(
        parent_widget, "Save JSON Export", "", "JSON Files (*.json)"
    )
    if not file_path:
        return

    nodes = database.get_nodes_for_project(project_id)
    coded_segments = database.get_coded_segments_for_project(project_id)

    segments_by_node = {}
    for seg in coded_segments:
        node_id = seg["node_id"]
        if node_id not in segments_by_node:
            segments_by_node[node_id] = []
        # Store a dictionary with participant and text
        segments_by_node[node_id].append(
            {
                "participant": seg["participant_name"] or "N/A",
                "text": seg["content_preview"],
                "document": seg["document_title"],
            }
        )

    def build_json_recursively(parent_id=None):
        children_data = []
        children = sorted(
            [n for n in nodes if n["parent_id"] == parent_id],
            key=lambda x: x["position"],
        )

        for node in children:
            node_obj = {
                "id": node["id"],
                "name": node["name"],
                "segments": segments_by_node.get(node["id"], []),
                "children": build_json_recursively(parent_id=node["id"]),
            }
            children_data.append(node_obj)
        return children_data

    json_output = build_json_recursively()

    # --- Save the JSON with error handling ---
    try:
        with open(file_path, "w", encoding="utf-8") as f:
            json.dump(json_output, f, ensure_ascii=False, indent=4)
        _show_export_saved(
            parent_widget, "JSON data successfully saved to:", file_path
        )
    except PermissionError:
        export_service.show_permission_error(parent_widget, file_path)
    except Exception as e:
        export_service.show_unexpected_error(parent_widget, "Export Error", e)


def export_project_to_excel_single_sheet(
    project_id, parent_widget=None, participant_id=None
):
    """Exports project coded segments to a single Excel worksheet."""
    _log_export_start("excel-single-sheet", project_id, participant_id=participant_id)
    file_path = _get_save_file_path(
        parent_widget,
        "Save Excel Report",
        _default_export_filename(
            f"NodeFlow_{_get_project_name(project_id)}_single_sheet_{_participant_file_suffix(project_id, participant_id)}",
            "xlsx",
        ),
        "Excel Files (*.xlsx)",
    )
    if not file_path:
        return

    nodes = database.get_nodes_for_project(project_id)
    coded_segments = _filter_segments_for_participant(
        database.get_coded_segments_for_project(project_id), participant_id
    )
    LOGGER.debug(
        "Building single-sheet workbook with %s nodes and %s coded segments",
        len(nodes),
        len(coded_segments),
    )
    node_names = {node["id"]: node["name"] for node in nodes}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(["Node", "Participant", "Coded Segment", "Document"])

    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font

    for seg in sorted(
        coded_segments,
        key=lambda segment: (
            node_names.get(segment["node_id"], ""),
            segment.get("participant_name") or "",
            segment.get("document_title") or "",
        ),
    ):
        ws.append(
            [
                node_names.get(seg["node_id"], seg["node_name"]),
                seg["participant_name"] or "N/A",
                seg["content_preview"],
                seg["document_title"],
            ]
        )

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 80
    ws.column_dimensions["D"].width = 40

    try:
        wb.save(file_path)
        _show_export_saved(
            parent_widget, "Excel report successfully saved to:", file_path
        )
    except PermissionError:
        export_service.show_permission_error(parent_widget, file_path)
    except Exception as e:
        export_service.show_unexpected_error(parent_widget, "Export Error", e)


def export_classification_workbook(
    project_id,
    parent_widget=None,
    participant_id=None,
    start_node_id=None,
    language="English",
):
    """Exports quotes plus a respondent-by-code matrix in one workbook."""
    _log_export_start(
        "excel-classification",
        project_id,
        participant_id=participant_id,
        node_id=start_node_id,
    )
    nodes = database.get_nodes_for_project(project_id)
    ordered_nodes, path_by_id = _build_scoped_node_paths(nodes, start_node_id)
    if not ordered_nodes:
        LOGGER.warning(
            "Classification export aborted: no nodes found in scope for node_id=%s",
            start_node_id,
        )
        return

    if start_node_id is None:
        filename_stem = (
            f"NodeFlow_{_get_project_name(project_id)}_classification_workbook_"
            f"{_participant_file_suffix(project_id, participant_id)}"
        )
    else:
        filename_stem = (
            f"NodeFlow_{_get_project_name(project_id)}_{ordered_nodes[0]['name']}_"
            f"classification_workbook_{_participant_file_suffix(project_id, participant_id)}"
        )

    file_path = _get_save_file_path(
        parent_widget,
        "Save Excel Report",
        _default_export_filename(filename_stem, "xlsx"),
        "Excel Files (*.xlsx)",
    )
    if not file_path:
        return

    node_ids_in_scope = [node["id"] for node in ordered_nodes]
    node_order = {node_id: index for index, node_id in enumerate(node_ids_in_scope)}
    node_path_labels = {
        node_id: " > ".join(path_by_id[node_id]) for node_id in node_ids_in_scope
    }

    segments = _filter_segments_for_participant(
        database.get_coded_segments_for_project(project_id), participant_id
    )
    scoped_segments = [
        segment for segment in segments if segment["node_id"] in node_order
    ]
    LOGGER.debug(
        "Building classification workbook with %s scoped nodes and %s scoped segments",
        len(node_ids_in_scope),
        len(scoped_segments),
    )
    scoped_segments.sort(
        key=lambda segment: (
            node_order[segment["node_id"]],
            _participant_display_name(segment),
            segment.get("document_title") or "",
            segment.get("segment_start") or 0,
            segment.get("segment_end") or 0,
            segment.get("id") or 0,
        )
    )

    wb = openpyxl.Workbook()
    quotes_ws = wb.active
    quotes_ws.title = _sanitize_sheet_name(
        get_translation("export.classification_quotes_sheet", language)
    )
    _set_header_row(
        quotes_ws,
        [
            get_translation("export.findings_column", language),
            get_translation("export.node_column", language),
            get_translation("export.subnode_column", language),
            get_translation("export.respondent_column", language),
            get_translation("export.original_quote_column", language),
            get_translation("export.document_column", language),
        ],
    )
    quotes_ws.freeze_panes = "A2"

    for segment in scoped_segments:
        findings, node, sub_node = _build_classification_columns(
            path_by_id[segment["node_id"]]
        )
        quotes_ws.append(
            [
                findings,
                node,
                sub_node,
                _participant_display_name(segment),
                segment["content_preview"],
                segment["document_title"],
            ]
        )

    for row in quotes_ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                horizontal="left", vertical="top", wrap_text=True
            )
    _set_column_widths(
        quotes_ws,
        {"A": 24, "B": 24, "C": 36, "D": 20, "E": 80, "F": 32},
    )

    matrix_ws = wb.create_sheet(
        title=_sanitize_sheet_name(
            get_translation("export.respondent_matrix_sheet", language)
        )
    )
    _set_header_row(
        matrix_ws,
        [get_translation("export.respondent_column", language)]
        + [node_path_labels[node_id] for node_id in node_ids_in_scope],
    )
    matrix_ws.freeze_panes = "B2"

    counts = Counter()
    participant_names = set()
    for segment in scoped_segments:
        participant_name = _participant_display_name(segment)
        participant_names.add(participant_name)
        counts[(participant_name, segment["node_id"])] += 1

    for participant_name in sorted(participant_names, key=str.casefold):
        matrix_ws.append(
            [participant_name]
            + [counts[(participant_name, node_id)] for node_id in node_ids_in_scope]
        )

    _set_column_widths(matrix_ws, {"A": 20})
    for row in matrix_ws.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="top")
    for index, node_id in enumerate(node_ids_in_scope, start=2):
        matrix_ws.column_dimensions[
            openpyxl.utils.cell.get_column_letter(index)
        ].width = max(18, min(36, len(node_path_labels[node_id]) + 4))

    _save_workbook_with_feedback(
        wb,
        file_path,
        parent_widget,
        "Excel report successfully saved to:",
    )


def export_to_excel(project_id, parent_widget=None, participant_id=None):
    """Exports coded segments to an Excel file with one sheet per node."""
    _log_export_start("excel-multi-sheet", project_id, participant_id=participant_id)

    file_path = _get_save_file_path(
        parent_widget,
        "Save Excel Report",
        _default_export_filename(
            f"NodeFlow_{_get_project_name(project_id)}_project_multi_sheet_{_participant_file_suffix(project_id, participant_id)}",
            "xlsx",
        ),
        "Excel Files (*.xlsx)",
    )
    if not file_path:
        return

    # Fetch data for the ENTIRE project
    nodes = database.get_nodes_for_project(project_id)
    coded_segments = _filter_segments_for_participant(
        database.get_coded_segments_for_project(project_id), participant_id
    )
    LOGGER.debug(
        "Building multi-sheet workbook with %s nodes and %s coded segments",
        len(nodes),
        len(coded_segments),
    )

    # --- Data Structuring for Traversal ---
    nodes_by_id = {n["id"]: n for n in nodes}
    nodes_by_parent = {n_id: [] for n_id in nodes_by_id}
    nodes_by_parent[None] = []  # For root nodes

    for n_id, node in nodes_by_id.items():
        nodes_by_parent.get(node["parent_id"], []).append(node)

    # Sort all children lists by position
    for children_list in nodes_by_parent.values():
        children_list.sort(key=lambda x: x["position"])

    # --- Descendant and Segment Mapping ---
    descendant_cache = {}

    def get_all_descendant_ids(node_id):
        if node_id in descendant_cache:
            return descendant_cache[node_id]
        descendants = []
        for child_node in nodes_by_parent.get(node_id, []):
            descendants.append(child_node["id"])
            descendants.extend(get_all_descendant_ids(child_node["id"]))
        descendant_cache[node_id] = descendants
        return descendants

    # --- Workbook Creation ---
    wb = openpyxl.Workbook()
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    def sanitize_sheet_name(name):
        name = re.sub(r"[\\/*?:\[\]]", "", name)
        return name[:31]

    # --- Recursive Sheet Creation ---
    def create_sheets_recursively(parent_id, prefix=""):
        children = nodes_by_parent.get(parent_id, [])
        for i, node in enumerate(children):
            node_id = node["id"]
            current_prefix = f"{prefix}{i + 1}."
            sheet_name = sanitize_sheet_name(f"{current_prefix} {node['name']}")

            # Create worksheet
            ws = wb.create_sheet(title=sheet_name)
            headers = ["Participant", "Coded Segment", "Document"]
            ws.append(headers)

            # --- NEW: Bold headers ---
            header_font = Font(bold=True)
            for cell in ws[1]:
                cell.font = header_font

            # Get all node IDs for this sheet (self + descendants)
            ids_to_include = [node_id] + get_all_descendant_ids(node_id)

            # Filter segments and write to sheet
            segments_for_sheet = [
                s for s in coded_segments if s["node_id"] in ids_to_include
            ]
            for seg in segments_for_sheet:
                participant = seg["participant_name"] or "N/A"
                ws.append([participant, seg["content_preview"], seg["document_title"]])

            # Adjust column widths
            ws.column_dimensions["A"].width = 25
            ws.column_dimensions["B"].width = 80
            ws.column_dimensions["C"].width = 40

            # Recurse for children
            create_sheets_recursively(node_id, prefix=current_prefix)

    # Start the process from the root (nodes with no parent)
    create_sheets_recursively(None)

    # --- Save the workbook with error handling ---
    try:
        wb.save(file_path)
        _show_export_saved(
            parent_widget, "Excel report successfully saved to:", file_path
        )
    except PermissionError:
        export_service.show_permission_error(parent_widget, file_path)
    except Exception as e:
        export_service.show_unexpected_error(parent_widget, "Export Error", e)


def export_node_family_to_word(project_id, start_node_id, parent_widget=None):
    """Exports a specific node and its children to a .docx file."""
    if not start_node_id:
        return

    # Fetch all project data
    all_nodes = database.get_nodes_for_project(project_id)
    coded_segments = database.get_coded_segments_for_project(project_id)
    nodes_map = {n["id"]: n for n in all_nodes}

    start_node = nodes_map.get(start_node_id)
    if not start_node:
        return

    file_path, _ = QFileDialog.getSaveFileName(
        parent_widget,
        f"Save Word Report for '{start_node['name']}'",
        "",
        "Word Documents (*.docx)",
    )
    if not file_path:
        return

    # Filter segments for all nodes in the family first
    segments_by_node = {}
    ids_to_include = [start_node_id] + get_all_descendant_ids(
        start_node_id, nodes_map, all_nodes
    )
    for seg in coded_segments:
        if seg["node_id"] in ids_to_include:
            node_id = seg["node_id"]
            if node_id not in segments_by_node:
                segments_by_node[node_id] = []
            segments_by_node[node_id].append(seg)

    # Build the document
    doc = Document()
    doc.add_heading(f"Report for Node: {start_node['name']}", 0)

    # --- REFACTORED: Correct recursive function ---
    def write_nodes_recursively(node_id, level, prefix):
        node = nodes_map.get(node_id)
        if not node:
            return

        # 1. Write the current node's info
        doc.add_heading(f"{prefix} {node['name']}", level=level)
        if node_id in segments_by_node:
            for seg_data in segments_by_node[node_id]:
                participant = seg_data["participant_name"] or "N/A"
                text = seg_data["content_preview"]
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"{participant}: ").bold = True
                p.add_run(text)

        # 2. Recurse for children
        children = sorted(
            [n for n in all_nodes if n["parent_id"] == node_id],
            key=lambda x: x["position"],
        )
        for i, child_node in enumerate(children):
            child_prefix = f"{prefix}{i + 1}."
            write_nodes_recursively(child_node["id"], level + 1, child_prefix)

    # Start the recursion with the selected node
    write_nodes_recursively(start_node_id, level=1, prefix="1.")

    # --- Save the document with error handling ---
    try:
        doc.save(file_path)
        # Show a message box with an 'Open File' button
        msg_box = QMessageBox(parent_widget)
        msg_box.setWindowTitle("Export Successful")
        msg_box.setText(f"Report successfully saved to:\n{file_path}")
        open_button = msg_box.addButton("Open File", QMessageBox.ActionRole)
        msg_box.addButton(QMessageBox.Ok)
        msg_box.exec_()
        if msg_box.clickedButton() == open_button:
            import os
            import sys

            if sys.platform.startswith("win"):
                os.startfile(file_path)
            elif sys.platform.startswith("darwin"):
                os.system(f'open "{file_path}"')
            else:
                os.system(f'xdg-open "{file_path}"')
    except PermissionError:
        QMessageBox.critical(
            parent_widget,
            "Permission Denied",
            f"Could not save the file to:\n{file_path}\n\nPlease make sure you have permissions to write to this location and that the file is not currently open in another program.",
        )
    except Exception as e:
        QMessageBox.critical(
            parent_widget,
            "Export Error",
            f"An unexpected error occurred while saving the file:\n{e}",
        )


def get_all_descendant_ids(start_node_id, nodes_map, all_nodes):
    """Helper to get all descendant IDs for a given node."""
    nodes_by_parent = {n_id: [] for n_id in nodes_map}
    nodes_by_parent[None] = []
    for n in all_nodes:
        nodes_by_parent.setdefault(n["parent_id"], []).append(n)

    descendants = []
    nodes_to_visit = [start_node_id]
    while nodes_to_visit:
        current_id = nodes_to_visit.pop(0)
        children = nodes_by_parent.get(current_id, [])
        for child in children:
            descendants.append(child["id"])
            nodes_to_visit.append(child["id"])
    return descendants


# --- NEW: Selective node family export to Excel ---
def export_node_family_to_excel(
    project_id, start_node_id, parent_widget=None, participant_id=None
):
    """Exports a specific node and its children to an .xlsx file."""
    _log_export_start(
        "excel-node-family-single-sheet",
        project_id,
        participant_id=participant_id,
        node_id=start_node_id,
    )
    if not start_node_id:
        return

    all_nodes = database.get_nodes_for_project(project_id)
    coded_segments = database.get_coded_segments_for_project(project_id)
    nodes_map = {n["id"]: n for n in all_nodes}

    start_node = nodes_map.get(start_node_id)
    if not start_node:
        LOGGER.warning("Node-family single-sheet export aborted: node %s not found", start_node_id)
        return

    file_path = _get_save_file_path(
        parent_widget,
        f"Save Excel Report for '{start_node['name']}'",
        _default_export_filename(
            f"NodeFlow_{_get_project_name(project_id)}_{start_node['name']}_single_sheet_{_participant_file_suffix(project_id, participant_id)}",
            "xlsx",
        ),
        "Excel Files (*.xlsx)",
    )
    if not file_path:
        return

    ids_to_include = [start_node_id] + get_all_descendant_ids(
        start_node_id, nodes_map, all_nodes
    )
    coded_segments = _filter_segments_for_participant(coded_segments, participant_id)
    LOGGER.debug(
        "Building node-family single-sheet workbook with %s nodes in scope and %s candidate segments",
        len(ids_to_include),
        len(coded_segments),
    )

    wb = openpyxl.Workbook()
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    def sanitize_sheet_name(name):
        return re.sub(r"[\\/*?:\[\]]", "", name)[:31]

    # Create one sheet for the whole family
    sheet_name = sanitize_sheet_name(start_node["name"])
    ws = wb.create_sheet(title=sheet_name)
    headers = ["Node", "Participant", "Coded Segment", "Document"]
    ws.append(headers)

    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font

    # Filter segments and write to sheet
    segments_for_sheet = [s for s in coded_segments if s["node_id"] in ids_to_include]
    for seg in sorted(segments_for_sheet, key=lambda s: s["node_name"]):
        ws.append(
            [
                seg["node_name"],
                seg["participant_name"] or "N/A",
                seg["content_preview"],
                seg["document_title"],
            ]
        )

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 80
    ws.column_dimensions["D"].width = 40

    # --- Save the workbook with error handling ---
    try:
        wb.save(file_path)
        # Show a message box with an 'Open File' button
        msg_box = QMessageBox(parent_widget)
        msg_box.setWindowTitle("Export Successful")
        msg_box.setText(f"Excel report successfully saved to:\n{file_path}")
        open_button = msg_box.addButton("Open File", QMessageBox.ActionRole)
        msg_box.addButton(QMessageBox.Ok)
        msg_box.exec_()
        if msg_box.clickedButton() == open_button:
            import os
            import sys

            if sys.platform.startswith("win"):
                os.startfile(file_path)
            elif sys.platform.startswith("darwin"):
                os.system(f'open "{file_path}"')
            else:
                os.system(f'xdg-open "{file_path}"')
    except PermissionError:
        QMessageBox.critical(
            parent_widget,
            "Permission Denied",
            f"Could not save the file to:\n{file_path}\n\nPlease make sure you have permissions to write to this location and that the file is not currently open in another program.",
        )
    except Exception as e:
        QMessageBox.critical(
            parent_widget,
            "Export Error",
            f"An unexpected error occurred while saving the file:\n{e}",
        )


def export_node_family_to_excel_multi_sheet(
    project_id, start_node_id, parent_widget=None, participant_id=None
):
    """Exports a specific node and its children to an .xlsx file with multiple sheets."""
    _log_export_start(
        "excel-node-family-multi-sheet",
        project_id,
        participant_id=participant_id,
        node_id=start_node_id,
    )
    if not start_node_id:
        return

    # Fetch all project data
    all_nodes = database.get_nodes_for_project(project_id)
    coded_segments = database.get_coded_segments_for_project(project_id)
    nodes_map = {n["id"]: n for n in all_nodes}

    start_node = nodes_map.get(start_node_id)
    if not start_node:
        LOGGER.warning("Node-family multi-sheet export aborted: node %s not found", start_node_id)
        return

    file_path = _get_save_file_path(
        parent_widget,
        f"Save Excel Report for '{start_node['name']}'",
        _default_export_filename(
            f"NodeFlow_{_get_project_name(project_id)}_{start_node['name']}_multi_sheet_{_participant_file_suffix(project_id, participant_id)}",
            "xlsx",
        ),
        "Excel Files (*.xlsx)",
    )
    if not file_path:
        return

    coded_segments = _filter_segments_for_participant(coded_segments, participant_id)
    LOGGER.debug(
        "Building node-family multi-sheet workbook with %s total nodes and %s candidate segments",
        len(all_nodes),
        len(coded_segments),
    )

    # --- Data Structuring for Traversal ---
    nodes_by_parent = {n_id: [] for n_id in nodes_map}
    nodes_by_parent[None] = []
    for n in all_nodes:
        nodes_by_parent.setdefault(n["parent_id"], []).append(n)

    for children_list in nodes_by_parent.values():
        children_list.sort(key=lambda x: x["position"])

    # --- Workbook Creation ---
    wb = openpyxl.Workbook()
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    def sanitize_sheet_name(name):
        return re.sub(r"[\\/*?:\[\]]", "", name)[:31]

    # --- Recursive Sheet Creation ---
    def create_sheets_recursively(node, prefix):
        node_id = node["id"]
        sheet_name = sanitize_sheet_name(f"{prefix} {node['name']}")

        ws = wb.create_sheet(title=sheet_name)
        headers = ["Participant", "Coded Segment", "Document"]
        ws.append(headers)

        header_font = Font(bold=True)
        for cell in ws[1]:
            cell.font = header_font

        ids_to_include_for_this_sheet = [node_id] + get_all_descendant_ids(
            node_id, nodes_map, all_nodes
        )

        segments_for_sheet = [
            s for s in coded_segments if s["node_id"] in ids_to_include_for_this_sheet
        ]
        for seg in segments_for_sheet:
            participant = seg["participant_name"] or "N/A"
            ws.append([participant, seg["content_preview"], seg["document_title"]])

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 80
        ws.column_dimensions["C"].width = 40

        # Recurse for children
        children = nodes_by_parent.get(node_id, [])
        for i, child_node in enumerate(children):
            create_sheets_recursively(child_node, f"{prefix}{i + 1}.")

    # Start the process from the start_node with prefix "1."
    create_sheets_recursively(start_node, "1.")

    # --- Save the workbook with error handling ---
    try:
        wb.save(file_path)
        # Show a message box with an 'Open File' button
        msg_box = QMessageBox(parent_widget)
        msg_box.setWindowTitle("Export Successful")
        msg_box.setText(f"Excel report successfully saved to:\n{file_path}")
        open_button = msg_box.addButton("Open File", QMessageBox.ActionRole)
        msg_box.addButton(QMessageBox.Ok)
        msg_box.exec_()
        if msg_box.clickedButton() == open_button:
            import os
            import sys

            if sys.platform.startswith("win"):
                os.startfile(file_path)
            elif sys.platform.startswith("darwin"):
                os.system(f'open "{file_path}"')
            else:
                os.system(f'xdg-open "{file_path}"')
    except PermissionError:
        QMessageBox.critical(
            parent_widget,
            "Permission Denied",
            f"Could not save the file to:\n{file_path}\n\nPlease make sure you have permissions to write to this location and that the file is not currently open in another program.",
        )
    except Exception as e:
        QMessageBox.critical(
            parent_widget,
            "Export Error",
            f"An unexpected error occurred while saving the file:\n{e}",
        )


def export_overall_participants_to_excel(project_id, parent_widget=None):
    """
    Exports a comprehensive Excel report with an overall node/participant view.
    Features include dynamic node level columns, a tree-view layout for nodes,
    specific cell formatting, and an option to open the file from the success dialog.
    """
    _log_export_start("excel-overall-participants", project_id)
    file_path = _get_save_file_path(
        parent_widget,
        "Save Overall Participants Report",
        _default_export_filename(
            f"NodeFlow_{_get_project_name(project_id)}_overall_participants",
            "xlsx",
        ),
        "Excel Files (*.xlsx)",
    )
    if not file_path:
        return

    try:
        # --- 1. Fetch all necessary data from the database ---
        nodes = database.get_nodes_for_project(project_id)
        participants = database.get_participants_for_project(project_id)
        all_segments = database.get_coded_segments_for_project(project_id)
        LOGGER.debug(
            "Building overall-participants workbook with %s nodes, %s participants, %s segments",
            len(nodes),
            len(participants),
            len(all_segments),
        )

        if not nodes:
            QMessageBox.information(
                parent_widget, "No Data", "There are no nodes to export."
            )
            return

        # --- 2. Structure the data for easy access ---
        participants_map = {p["id"]: p["name"] for p in participants}
        sorted_participant_ids = sorted(
            participants_map.keys(), key=lambda pid: participants_map[pid]
        )

        G = nx.DiGraph()
        nodes_by_parent = defaultdict(list)
        for node in nodes:
            nodes_by_parent[node["parent_id"]].append(node)
            if node["parent_id"] is not None:
                G.add_edge(node["parent_id"], node["id"])

        max_depth = 0
        root_nodes = [n for n in nodes if n["parent_id"] is None]
        for root_node in root_nodes:
            descendants = nx.descendants(G, root_node["id"])
            for descendant in descendants:
                try:
                    path_length = nx.shortest_path_length(
                        G, root_node["id"], descendant
                    )
                    max_depth = max(max_depth, path_length + 1)
                except nx.NetworkXNoPath:
                    pass
        if not max_depth and root_nodes:
            max_depth = 1

        for parent_id in nodes_by_parent:
            nodes_by_parent[parent_id].sort(key=lambda n: n["position"])

        segments_by_node_participant = defaultdict(lambda: defaultdict(list))
        for seg in all_segments:
            segments_by_node_participant[seg["node_id"]][seg["participant_id"]].append(
                seg["content_preview"]
            )

        # --- 3. Create the Workbook and Overall Sheet ---
        wb = openpyxl.Workbook()
        ws_overall = wb.active
        ws_overall.title = "Overall"

        headers = (
            [f"Level {i+1}" for i in range(max_depth)] if max_depth > 0 else []
        ) + [participants_map[pid] for pid in sorted_participant_ids]
        ws_overall.append(headers)

        for cell in ws_overall[1]:
            cell.alignment = Alignment(horizontal="left", vertical="top")
            cell.font = Font(bold=True)

        # --- 4. Recursively write data to the Overall sheet ---
        def write_rows_recursively(parent_id, level):
            for node in nodes_by_parent.get(parent_id, []):
                row_data = [""] * max_depth
                if (level - 1) < max_depth:
                    row_data[level - 1] = node["name"]

                participant_cells = []
                for pid in sorted_participant_ids:
                    segments = segments_by_node_participant[node["id"]].get(pid, [])
                    # Encapsulate each segment in quotes and prepend a hyphen
                    formatted_segments = [f'- "{s}"' for s in segments]
                    participant_cells.append("\n".join(formatted_segments))

                ws_overall.append(row_data + participant_cells)

                current_row_index = ws_overall.max_row
                for col_idx in range(1, len(row_data) + len(participant_cells) + 1):
                    cell = ws_overall.cell(row=current_row_index, column=col_idx)
                    cell.alignment = Alignment(
                        wrap_text=True, vertical="top", horizontal="left"
                    )

                write_rows_recursively(node["id"], level + 1)

        write_rows_recursively(None, 1)

        # --- 5. Adjust column widths ---
        for i in range(max_depth):
            col_letter = openpyxl.utils.cell.get_column_letter(i + 1)
            ws_overall.column_dimensions[col_letter].width = 30

        for i in range(len(sorted_participant_ids)):
            col_letter = openpyxl.utils.cell.get_column_letter(max_depth + 1 + i)
            ws_overall.column_dimensions[col_letter].width = 50

        # --- 6. Save the workbook ---
        wb.save(file_path)

        _show_export_saved(parent_widget, "The report has been saved to:", file_path)

    except PermissionError:
        export_service.show_permission_error(parent_widget, file_path)
    except Exception as e:
        export_service.show_unexpected_error(parent_widget, "Export Error", e)


def export_co_occurrence_to_gexf(project_id, parent_widget=None):
    """Exports the code co-occurrence data to a GEXF file for network analysis."""
    file_path, _ = QFileDialog.getSaveFileName(
        parent_widget, "Save GEXF File", "", "GEXF Files (*.gexf)"
    )
    if not file_path:
        return

    try:
        nodes = database.get_nodes_for_project(project_id)
        segments = database.get_coded_segments_for_project(project_id)

        node_map = {node["id"]: node["name"] for node in nodes}
        co_occurrence_matrix = {}

        segments_by_content = {}
        for seg in segments:
            seg_key = (
                seg["document_id"],
                seg["segment_start"],
                seg["segment_end"],
            )
            if seg_key not in segments_by_content:
                segments_by_content[seg_key] = []
            segments_by_content[seg_key].append(seg["node_id"])

        for seg_key, node_ids in segments_by_content.items():
            unique_node_ids = sorted(list(set(node_ids)))
            for i in range(len(unique_node_ids)):
                for j in range(i, len(unique_node_ids)):
                    node1_id = unique_node_ids[i]
                    node2_id = unique_node_ids[j]

                    if node1_id == node2_id:
                        continue

                    node1_name = node_map.get(node1_id)
                    node2_name = node_map.get(node2_id)

                    if not node1_name or not node2_name:
                        continue

                    co_occurrence_matrix.setdefault(node1_name, {})
                    co_occurrence_matrix[node1_name].setdefault(node2_name, 0)
                    co_occurrence_matrix[node1_name][node2_name] += 1

                    co_occurrence_matrix.setdefault(node2_name, {})
                    co_occurrence_matrix[node2_name].setdefault(node1_name, 0)
                    co_occurrence_matrix[node2_name][node1_name] += 1

        G = nx.Graph()
        for node1, connections in co_occurrence_matrix.items():
            for node2, weight in connections.items():
                if weight > 0 and node1 != node2:
                    G.add_edge(node1, node2, weight=weight)

        if not G.nodes():
            for node in nodes:
                G.add_node(node["name"])

        nx.write_gexf(G, file_path)

        # Show a message box with an 'Open File' button
        msg_box = QMessageBox(parent_widget)
        msg_box.setWindowTitle("Export Successful")
        msg_box.setText(f"GEXF file successfully saved to:\n{file_path}")
        open_button = msg_box.addButton("Open File", QMessageBox.ActionRole)
        msg_box.addButton(QMessageBox.Ok)
        msg_box.exec_()
        if msg_box.clickedButton() == open_button:
            import os
            import sys

            if sys.platform.startswith("win"):
                os.startfile(file_path)
            elif sys.platform.startswith("darwin"):
                os.system(f'open "{file_path}"')
            else:
                os.system(f'xdg-open "{file_path}"')
    except Exception as e:
        LOGGER.exception("GEXF export failed")
        export_service.show_unexpected_error(parent_widget, "Export Error", e)


def export_annotated_document(
    project_id, document_id, document_title, parent_widget=None
):
    """Exports a single document with its coded segments highlighted."""
    if not document_id:
        QMessageBox.warning(
            parent_widget, "No Document", "No document is currently loaded."
        )
        return

    file_path, _ = QFileDialog.getSaveFileName(
        parent_widget,
        f"Save Annotated Document '{document_title}'",
        "",
        "Word Documents (*.docx)",
    )
    if not file_path:
        return

    try:
        content, _ = database.get_document_content(document_id)
        segments = database.get_coded_segments_for_document(document_id)
        segments.sort(key=lambda s: s["segment_start"])

        doc = Document()
        doc.add_heading(f"Annotated Document: {document_title}", 0)

        p = doc.add_paragraph()
        last_pos = 0
        for seg in segments:
            # Add text before the segment
            p.add_run(content[last_pos : seg["segment_start"]])

            # Add the segment with background color for highlight
            run = p.add_run(seg["content_preview"])
            hex_color = seg["node_color"].lstrip("#")
            rgb = tuple(int(hex_color[i : i + 2], 16) for i in (0, 2, 4))
            # Set background color using shading (python-docx)
            run.font.highlight_color = None
            rPr = run._element.get_or_add_rPr()
            shd = rPr.xpath("./w:shd")
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement

            if not shd:
                shd_elem = OxmlElement("w:shd")
                shd_elem.set(qn("w:fill"), f"{hex_color}")
                shd_elem.set(qn("w:val"), "clear")
                shd_elem.set(qn("w:color"), "auto")
                rPr.append(shd_elem)
            else:
                shd[0].set(qn("w:fill"), f"{hex_color}")
                shd[0].set(qn("w:val"), "clear")
                shd[0].set(qn("w:color"), "auto")
            # Set font color for contrast
            brightness = (rgb[0] * 299 + rgb[1] * 587 + rgb[2] * 114) / 1000
            if brightness > 128:
                run.font.color.rgb = RGBColor(0, 0, 0)  # black
            else:
                run.font.color.rgb = RGBColor(255, 255, 255)  # white

            # Add remark as [<node>] (no field name prefix)
            info_run = p.add_run(f" [{seg['node_name']}] ")
            info_run.italic = True
            info_run.font.size = run.font.size

            last_pos = seg["segment_end"]

        # Add remaining text
        p.add_run(content[last_pos:])

        doc.save(file_path)

        # Show a message box with an 'Open File' button
        _show_export_saved(
            parent_widget,
            "Annotated document successfully saved to:",
            file_path,
        )
    except Exception as e:
        LOGGER.exception("Annotated document export failed")
        export_service.show_unexpected_error(parent_widget, "Export Error", e)
