from __future__ import annotations

import hashlib
import shutil
from datetime import datetime
from pathlib import Path

import docx
import openpyxl

import database
from utils.app_paths import get_imports_dir


class ImportService:
    def read_text_document(self, file_path: str) -> tuple[str, str, dict]:
        path = Path(file_path)
        if path.suffix.lower() == ".docx":
            document = docx.Document(path)
            content = "\n\n".join(paragraph.text for paragraph in document.paragraphs)
        else:
            content = path.read_text(encoding="utf-8")
        metadata = self.copy_source_metadata(path, path.suffix.lower().lstrip("."))
        return path.name, content, metadata

    def copy_source_metadata(
        self, source_path: Path | str, source_kind: str, source_row: int | None = None
    ) -> dict:
        source_path = Path(source_path)
        copied_path, digest = self._copy_import_source(source_path)
        return {
            "source_filename": source_path.name,
            "source_copy_path": str(copied_path),
            "source_sha256": digest,
            "source_kind": source_kind,
            "source_row": source_row,
            "imported_at": datetime.now().isoformat(timespec="seconds"),
        }

    @staticmethod
    def _copy_import_source(source_path: Path) -> tuple[Path, str]:
        digest = hashlib.sha256()
        with source_path.open("rb") as source:
            for chunk in iter(lambda: source.read(1024 * 1024), b""):
                digest.update(chunk)
        source_hash = digest.hexdigest()
        suffix = source_path.suffix
        safe_stem = "".join(
            char if char.isalnum() or char in ("-", "_") else "_"
            for char in source_path.stem
        ).strip("_") or "import"
        target = get_imports_dir() / f"{safe_stem}_{source_hash[:12]}{suffix}"
        if not target.exists():
            shutil.copy2(source_path, target)
        return target, source_hash

    def import_excel_data(
        self,
        project_id: int,
        file_path: str,
        mappings: dict[str, str],
    ) -> tuple[int, list[str]]:
        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True)
            sheet = workbook.active
        except Exception as exc:
            return 0, [f"Failed to open or read the Excel file: {exc}"]
        workbook_metadata = self.copy_source_metadata(file_path, "xlsx")

        headers = [cell.value for cell in sheet[1]]
        try:
            title_col_idx = headers.index(mappings["title"])
            content_col_idx = headers.index(mappings["content"])
            participant_col_name = mappings.get("participant")
            participant_col_idx = (
                headers.index(participant_col_name)
                if participant_col_name and participant_col_name != "<Assign Later>"
                else None
            )
        except ValueError as exc:
            return 0, [f"A mapped column was not found in the Excel file: {exc}"]

        participants_in_db = {
            p["name"]: p["id"] for p in database.get_participants_for_project(project_id)
        }
        existing_titles = {
            doc["title"] for doc in database.get_documents_for_project(project_id)
        }

        docs_imported = 0
        errors: list[str] = []

        for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            title = row[title_col_idx].value
            content = row[content_col_idx].value
            if not title or not content:
                errors.append(f"Row {row_idx}: Skipped due to empty title or content.")
                continue

            title = str(title).strip()
            if title in existing_titles:
                base_title = title
                counter = 1
                while title in existing_titles:
                    title = f"{base_title} (copy {counter})"
                    counter += 1

            content = str(content).strip()
            participant_id = None
            if participant_col_idx is not None:
                participant_name = row[participant_col_idx].value
                if participant_name:
                    participant_name = str(participant_name).strip()
                    participant_id = participants_in_db.get(participant_name)
                    if participant_id is None:
                        try:
                            database.add_participant(project_id, participant_name)
                            participant_id = next(
                                (
                                    p["id"]
                                    for p in database.get_participants_for_project(project_id)
                                    if p["name"] == participant_name
                                ),
                                None,
                            )
                            participants_in_db[participant_name] = participant_id
                        except Exception as exc:
                            errors.append(
                                f"Row {row_idx}: Could not create participant '{participant_name}'. Error: {exc}"
                            )
                            continue

            try:
                source_metadata = dict(workbook_metadata)
                source_metadata["source_row"] = row_idx
                database.add_document(
                    project_id, title, content, participant_id, source_metadata
                )
                existing_titles.add(title)
                docs_imported += 1
            except Exception as exc:
                errors.append(
                    f"Row {row_idx}: Failed to import document '{title}'. Error: {exc}"
                )

        return docs_imported, errors


import_service = ImportService()
